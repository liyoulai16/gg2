import csv
import os
import shutil
from datetime import datetime
from typing import List, Dict, Optional, Tuple
import sqlite3


try:
    import openpyxl
    from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
    from openpyxl.utils import get_column_letter
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


class ImportExportManager:
    def __init__(self, db):
        self.db = db

    def export_transactions_to_csv(self, filepath: str, transactions: List[Dict] = None,
                                     start_date: str = None, end_date: str = None,
                                     account_id: int = None, type_: str = None) -> Tuple[bool, str]:
        try:
            if transactions is None:
                transactions = self.db.get_transactions(
                    account_id=account_id,
                    start_date=start_date,
                    end_date=end_date,
                    type_=type_
                )

            if not transactions:
                return False, "没有可导出的交易记录"

            with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
                writer = csv.writer(f)
                headers = ['ID', '日期', '类型', '账户', '分类', '金额', '描述', '创建时间']
                writer.writerow(headers)

                for t in transactions:
                    type_label = '收入' if t['type'] == 'income' else '支出'
                    row = [
                        t['id'],
                        t['date'],
                        type_label,
                        t['account_name'],
                        t['category_name'],
                        t['amount'],
                        t['description'] or '',
                        t['created_at']
                    ]
                    writer.writerow(row)

            return True, f"成功导出 {len(transactions)} 条记录到 {filepath}"
        except Exception as e:
            return False, f"导出失败: {str(e)}"

    def export_transactions_to_excel(self, filepath: str, transactions: List[Dict] = None,
                                       start_date: str = None, end_date: str = None,
                                       account_id: int = None, type_: str = None) -> Tuple[bool, str]:
        if not HAS_OPENPYXL:
            return False, "未安装 openpyxl 库，无法导出 Excel。请运行: pip install openpyxl"

        try:
            if transactions is None:
                transactions = self.db.get_transactions(
                    account_id=account_id,
                    start_date=start_date,
                    end_date=end_date,
                    type_=type_
                )

            if not transactions:
                return False, "没有可导出的交易记录"

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "交易记录"

            headers = ['ID', '日期', '类型', '账户', '分类', '金额', '描述', '创建时间']
            header_fill = PatternFill(start_color="4285f4", end_color="4285f4", fill_type="solid")
            header_font = Font(bold=True, color="FFFFFF")
            header_alignment = Alignment(horizontal="center", vertical="center")

            for col_idx, header in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col_idx, value=header)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = header_alignment

            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin')
            )

            for row_idx, t in enumerate(transactions, 2):
                type_label = '收入' if t['type'] == 'income' else '支出'
                row_data = [
                    t['id'],
                    t['date'],
                    type_label,
                    t['account_name'],
                    t['category_name'],
                    t['amount'],
                    t['description'] or '',
                    t['created_at']
                ]

                for col_idx, value in enumerate(row_data, 1):
                    cell = ws.cell(row=row_idx, column=col_idx, value=value)
                    cell.border = thin_border
                    cell.alignment = Alignment(vertical="center")

                    if col_idx == 6:
                        cell.number_format = '#,##0.00'

            for col_idx in range(1, len(headers) + 1):
                column_letter = get_column_letter(col_idx)
                max_length = len(str(headers[col_idx - 1]))
                for row_idx in range(1, len(transactions) + 2):
                    try:
                        cell_value = str(ws.cell(row=row_idx, column=col_idx).value)
                        max_length = max(max_length, len(cell_value))
                    except:
                        pass
                adjusted_width = min(max_length + 2, 50)
                ws.column_dimensions[column_letter].width = adjusted_width

            wb.save(filepath)
            return True, f"成功导出 {len(transactions)} 条记录到 {filepath}"
        except Exception as e:
            return False, f"导出失败: {str(e)}"

    def backup_database(self, backup_dir: str = None) -> Tuple[bool, str]:
        try:
            db_path = self.db.db_path
            if not os.path.exists(db_path):
                return False, "数据库文件不存在"

            if backup_dir is None:
                backup_dir = os.path.dirname(db_path)

            os.makedirs(backup_dir, exist_ok=True)

            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            backup_filename = f"accounting_backup_{timestamp}.db"
            backup_path = os.path.join(backup_dir, backup_filename)

            shutil.copy2(db_path, backup_path)

            return True, f"数据库备份成功: {backup_path}"
        except Exception as e:
            return False, f"备份失败: {str(e)}"

    def import_transactions_from_csv(self, filepath: str, 
                                       account_map: Dict[str, int] = None,
                                       category_map: Dict[Tuple[str, str], int] = None,
                                       create_missing: bool = False) -> Tuple[bool, str, List[Dict]]:
        try:
            if not os.path.exists(filepath):
                return False, "文件不存在", []

            accounts = {a['name']: a for a in self.db.get_all_accounts()}
            categories = {(c['name'], c['type']): c for c in self.db.get_all_categories()}

            transactions = []
            errors = []
            success_count = 0

            with open(filepath, 'r', encoding='utf-8-sig') as f:
                reader = csv.DictReader(f)

                for row_idx, row in enumerate(reader, 2):
                    try:
                        date = row.get('日期', '').strip()
                        type_str = row.get('类型', '').strip()
                        account_name = row.get('账户', '').strip()
                        category_name = row.get('分类', '').strip()
                        amount_str = row.get('金额', '').strip()
                        description = row.get('描述', '').strip()

                        if not all([date, type_str, account_name, category_name, amount_str]):
                            errors.append(f"第 {row_idx} 行: 缺少必要字段")
                            continue

                        type_ = 'income' if type_str in ['收入', 'income'] else 'expense'

                        try:
                            amount = float(amount_str.replace(',', ''))
                            if amount <= 0:
                                raise ValueError("金额必须大于0")
                        except ValueError:
                            errors.append(f"第 {row_idx} 行: 无效的金额 '{amount_str}'")
                            continue

                        if account_name not in accounts:
                            if create_missing:
                                if self.db.add_account(account_name, 0):
                                    accounts = {a['name']: a for a in self.db.get_all_accounts()}
                                else:
                                    errors.append(f"第 {row_idx} 行: 无法创建账户 '{account_name}'")
                                    continue
                            else:
                                errors.append(f"第 {row_idx} 行: 账户 '{account_name}' 不存在")
                                continue

                        account_id = accounts[account_name]['id']

                        category_key = (category_name, type_)
                        if category_key not in categories:
                            if create_missing:
                                if self.db.add_category(category_name, type_):
                                    categories = {(c['name'], c['type']): c for c in self.db.get_all_categories()}
                                else:
                                    errors.append(f"第 {row_idx} 行: 无法创建分类 '{category_name}'")
                                    continue
                            else:
                                errors.append(f"第 {row_idx} 行: 分类 '{category_name}' 不存在")
                                continue

                        category_id = categories[category_key]['id']

                        if self.db.add_transaction(
                            account_id=account_id,
                            category_id=category_id,
                            type_=type_,
                            amount=amount,
                            description=description,
                            date=date
                        ):
                            success_count += 1
                            transactions.append({
                                'row': row_idx,
                                'date': date,
                                'type': type_,
                                'account': account_name,
                                'category': category_name,
                                'amount': amount,
                                'description': description
                            })
                        else:
                            errors.append(f"第 {row_idx} 行: 添加交易记录失败")

                    except Exception as e:
                        errors.append(f"第 {row_idx} 行: 处理错误 - {str(e)}")

            result_msg = f"导入完成: 成功 {success_count} 条"
            if errors:
                result_msg += f"，失败 {len(errors)} 条"

            return (success_count > 0), result_msg, errors

        except Exception as e:
            return False, f"导入失败: {str(e)}", [str(e)]

    def export_all_data_to_csv(self, directory: str) -> Tuple[bool, str]:
        try:
            os.makedirs(directory, exist_ok=True)
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')

            accounts = self.db.get_all_accounts()
            if accounts:
                filepath = os.path.join(directory, f"accounts_{timestamp}.csv")
                with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
                    writer = csv.writer(f)
                    writer.writerow(['ID', '账户名称', '余额', '创建时间'])
                    for a in accounts:
                        writer.writerow([a['id'], a['name'], a['balance'], a['created_at']])

            categories = self.db.get_all_categories()
            if categories:
                filepath = os.path.join(directory, f"categories_{timestamp}.csv")
                with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
                    writer = csv.writer(f)
                    writer.writerow(['ID', '分类名称', '类型', '创建时间'])
                    for c in categories:
                        type_label = '收入' if c['type'] == 'income' else '支出'
                        writer.writerow([c['id'], c['name'], type_label, c['created_at']])

            transactions = self.db.get_transactions()
            if transactions:
                filepath = os.path.join(directory, f"transactions_{timestamp}.csv")
                self.export_transactions_to_csv(filepath, transactions)

            budgets = self.db.get_all_budgets()
            if budgets:
                filepath = os.path.join(directory, f"budgets_{timestamp}.csv")
                with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
                    writer = csv.writer(f)
                    writer.writerow(['ID', '年份', '月份', '预算金额', '创建时间', '更新时间'])
                    for b in budgets:
                        writer.writerow([b['id'], b['year'], b['month'], b['amount'], b['created_at'], b['updated_at']])

            debts = self.db.get_all_debts()
            if debts:
                filepath = os.path.join(directory, f"debts_{timestamp}.csv")
                with open(filepath, 'w', newline='', encoding='utf-8-sig') as f:
                    writer = csv.writer(f)
                    writer.writerow(['ID', '类型', '对方', '金额', '剩余金额', '利率', 
                                     '开始日期', '到期日期', '状态', '描述', '创建时间', '更新时间'])
                    for d in debts:
                        type_label = '借出' if d['type'] == 'lend' else '借入'
                        writer.writerow([d['id'], type_label, d['counterparty'], d['amount'], 
                                        d['remaining_amount'], d['interest_rate'], d['start_date'],
                                        d['due_date'] or '', d['status'], d['description'] or '',
                                        d['created_at'], d['updated_at']])

            return True, f"成功导出所有数据到目录: {directory}"
        except Exception as e:
            return False, f"导出失败: {str(e)}"

    def restore_database(self, backup_path: str) -> Tuple[bool, str]:
        try:
            if not os.path.exists(backup_path):
                return False, "备份文件不存在"

            db_path = self.db.db_path
            temp_path = db_path + '.temp'

            try:
                conn = sqlite3.connect(backup_path)
                cursor = conn.cursor()
                cursor.execute("SELECT name FROM sqlite_master WHERE type='table'")
                tables = cursor.fetchall()
                conn.close()

                required_tables = ['accounts', 'categories', 'transactions']
                has_required = any(t[0] in required_tables for t in tables)
                if not has_required:
                    return False, "所选文件不是有效的记账应用数据库备份"
            except Exception as e:
                return False, f"无效的数据库文件: {str(e)}"

            if os.path.exists(db_path):
                shutil.copy2(db_path, temp_path)

            try:
                shutil.copy2(backup_path, db_path)
                if os.path.exists(temp_path):
                    os.remove(temp_path)
                return True, "数据库恢复成功，请重启应用以加载新数据"
            except Exception as e:
                if os.path.exists(temp_path):
                    shutil.copy2(temp_path, db_path)
                    os.remove(temp_path)
                return False, f"恢复失败: {str(e)}"

        except Exception as e:
            return False, f"恢复失败: {str(e)}"
